"""Text extraction. Optional engines (PyMuPDF, python-docx, pytesseract) are used when installed."""
import io
from dataclasses import dataclass

from .validation import UploadError


@dataclass
class Extracted:
    text: str
    pages: int
    kind: str  # extension the frontend should treat the text as (txt or csv)


def available() -> dict:
    out = {}
    for key, mod in {"pdf": "fitz", "docx": "docx", "ocr": "pytesseract"}.items():
        try:
            __import__(mod)
            out[key] = True
        except Exception:
            out[key] = False
    return out


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", errors="replace")


def extract(ext: str, data: bytes) -> Extracted:
    if ext in {"txt", "md", "log", "json"}:
        t = _decode(data)
        return Extracted(t, max(1, len(t) // 1800 + 1), "txt")
    if ext == "csv":
        t = _decode(data)
        return Extracted(t, max(1, len(t) // 1800 + 1), "csv")
    if ext == "pdf":
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise UploadError("PDF extraction is not installed on this server (install PyMuPDF).")
        try:
            doc = fitz.open(stream=data, filetype="pdf")
        except Exception:
            raise UploadError("This PDF is corrupt or encrypted and could not be opened.")
        pages = [p.get_text() for p in doc]
        text = "\n\n".join(pages).strip()
        if not text:
            raise UploadError("No text layer found. This looks like a scanned PDF; OCR is required.")
        return Extracted(text, len(pages), "txt")
    if ext == "docx":
        try:
            import docx
        except ImportError:
            raise UploadError("DOCX extraction is not installed on this server (install python-docx).")
        try:
            d = docx.Document(io.BytesIO(data))
        except Exception:
            raise UploadError("This DOCX file is corrupt and could not be opened.")
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for tb in d.tables:
            for row in tb.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts).strip()
        if not text:
            raise UploadError("No readable text was found in this document.")
        return Extracted(text, max(1, len(text) // 1800 + 1), "txt")
    if ext in {"png", "jpg", "jpeg", "tif"}:
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            raise UploadError("OCR is not installed on this server (install pytesseract and Tesseract).")
        try:
            text = pytesseract.image_to_string(Image.open(io.BytesIO(data))).strip()
        except Exception:
            raise UploadError("OCR failed on this image.")
        if not text:
            raise UploadError("OCR found no text in this image.")
        return Extracted(text, 1, "txt")
    raise UploadError("Unsupported file type.")
